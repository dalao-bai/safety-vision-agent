from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Callable

from docx import Document
from docx.shared import Pt

from app.persistence.models import Hazard


def build_docx_report(
    session_id: str,
    hazards: list[Hazard],
    report_dir: str,
    object_name_for: Callable[[str], str],
    hazard_name_for: Callable[[str | None], str],
    remediation_for: Callable[[str], list[dict[str, str]]],
    rule_blocks_for: Callable[[str, str | None], list[dict]] | None = None,
) -> str:
    """生成会话已确认隐患的.docx巡查报告，返回文件路径。"""
    out_dir = Path(report_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_heading("四口五临边隐患排查报告", 0)

    meta = doc.add_paragraph()
    meta.add_run(f"会话：{session_id}\n")
    meta.add_run(f"生成时间：{datetime.now(timezone.utc).isoformat(timespec='seconds')}\n")
    meta.add_run(f"已确认隐患数：{len(hazards)}")

    if not hazards:
        doc.add_paragraph("未发现已确认隐患。")
    else:
        for i, h in enumerate(hazards, 1):
            doc.add_heading(
                f"{i}. {object_name_for(h.object_id)} — {hazard_name_for(h.hazard_type_id)}",
                level=1,
            )
            # 若 VLM 未返回 rule_basis，从 KG rule_blocks 补全
            rule_basis = h.rule_basis
            if not rule_basis and rule_blocks_for:
                blocks = rule_blocks_for(h.object_id, h.hazard_type_id)
                if blocks:
                    rule_basis = "；".join(
                        f"{b['rule_text']}（{b['source']}）" if b.get("source") else b["rule_text"]
                        for b in blocks if b.get("rule_text")
                    )

            info = doc.add_paragraph()
            info.add_run(f"状态：{h.status}\n")
            info.add_run(f"位置 bbox：{h.bbox}\n")
            info.add_run(f"视觉证据：{h.visual_evidence}\n")
            info.add_run(f"规则依据：{rule_basis or '暂无'}")

            rem = remediation_for(h.object_id)
            if rem:
                doc.add_paragraph("整改建议（基于合格条件）：")
                for item in rem:
                    cond = item.get("condition", "")
                    src = f"（{item['source']}）" if item.get("source") else ""
                    doc.add_paragraph(f"{cond}{src}", style="List Bullet")

    path = out_dir / f"session_{session_id}_report.docx"
    doc.save(str(path))
    return str(path)
