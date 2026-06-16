# backend/app/services/foe_report_generator.py
"""四口五临边报告渲染 (.docx) + FOE 专属进程内任务注册表。

独立于 report_generator 的任务注册表，以保证下载链接指向 /api/foe/report/download。
"""

from __future__ import annotations

import os
import uuid
from typing import Any

from docx import Document
from docx.shared import Inches

from app.models.foe_schemas import FoeAnalysis

_STATUS_LABELS = {"confirmed_hazard": "确认隐患", "safe": "安全", "uncertain": "不确定"}
_STEP_LABELS = {"observe": "观察", "locate": "定位", "match_rule": "匹配规则", "assess": "评估"}
_IMAGE_WIDTH = Inches(3.5)


def generate_foe_report_docx(
    analyses: list[FoeAnalysis], retriever, output_dir: str, title: str | None = None
) -> str:
    """生成 .docx 报告，返回绝对路径。retriever 须有 retrieve(FoeObject)->[ClauseRef]。"""
    os.makedirs(output_dir, exist_ok=True)
    doc = Document()
    doc.add_heading(title or "四口五临边安全隐患报告", level=0)

    total_objs = sum(len(a.objects) for a in analyses)
    hazard_objs = sum(1 for a in analyses for o in a.objects if o.status == "confirmed_hazard")
    doc.add_paragraph(f"图片数：{len(analyses)}")
    doc.add_paragraph(f"识别对象数：{total_objs}")
    doc.add_paragraph(f"其中确认隐患：{hazard_objs}")

    for idx, analysis in enumerate(analyses, 1):
        doc.add_heading(f"图片 {idx}", level=1)
        ref_path = analysis.image_ref
        if ref_path and os.path.exists(ref_path):
            try:
                doc.add_picture(ref_path, width=_IMAGE_WIDTH)
            except Exception as exc:  # noqa: BLE001 - 缺图不应中断报告
                doc.add_paragraph(f"（图片无法嵌入：{ref_path}，{exc}）")

        if not analysis.objects:
            doc.add_paragraph("未识别到四口五临边对象。")
            continue

        for obj in analysis.objects:
            doc.add_heading(obj.related_object, level=2)
            doc.add_paragraph(f"状态：{_STATUS_LABELS.get(obj.status, obj.status)}")
            if obj.hazard_type:
                doc.add_paragraph(f"隐患类型：{obj.hazard_type}")
            doc.add_paragraph(f"边界框：{obj.object_bbox}")
            if obj.reasoning_chain:
                doc.add_paragraph("推理过程：")
                for step in obj.reasoning_chain:
                    label = _STEP_LABELS.get(step.step, step.step)
                    doc.add_paragraph(f"{label}：{step.content}")
            if obj.visual_evidence:
                doc.add_paragraph(f"视觉证据：{obj.visual_evidence}")
            if obj.evidence_sufficiency:
                doc.add_paragraph(f"证据充分性：{obj.evidence_sufficiency}")
            if obj.status == "uncertain" and obj.uncertainty_reason:
                doc.add_paragraph(f"不确定原因：{obj.uncertainty_reason}")

            refs = retriever.retrieve(obj)
            if refs:
                doc.add_paragraph("依据条款：")
                for ref in refs:
                    para = doc.add_paragraph()
                    para.add_run(f"【{ref.standard_code} {ref.clause_id}】").bold = True
                    para.add_run(ref.official_text)
            else:
                doc.add_paragraph("（未匹配到条款，建议人工复核）")

    filename = uuid.uuid4().hex + ".docx"
    out_path = os.path.abspath(os.path.join(output_dir, filename))
    doc.save(out_path)
    return out_path


# --- FOE 专属进程内任务注册表 -----------------------------------------------
_TASKS: dict[str, dict[str, Any]] = {}


def create_task(user_id: str) -> str:
    task_id = uuid.uuid4().hex
    _TASKS[task_id] = {"status": "pending", "user_id": user_id, "path": None, "error": None}
    return task_id


def set_task_done(task_id: str, path: str) -> None:
    task = _TASKS.get(task_id)
    if task is not None:
        task["status"] = "done"
        task["path"] = path


def set_task_error(task_id: str, msg: str) -> None:
    task = _TASKS.get(task_id)
    if task is not None:
        task["status"] = "error"
        task["error"] = msg


def get_task(task_id: str) -> dict[str, Any] | None:
    task = _TASKS.get(task_id)
    if task is None:
        return None
    view: dict[str, Any] = {"status": task["status"], "user_id": task["user_id"]}
    if task["status"] == "done":
        view["download_url"] = f"/api/foe/report/download/{task_id}"
        view["path"] = task["path"]
    elif task["status"] == "error":
        view["error"] = task["error"]
    return view
