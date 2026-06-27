from pathlib import Path

from docx import Document

from app.persistence.models import Hazard
from app.reports.builder import build_docx_report


def _doc_text(path: str) -> str:
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs)


def _hazards():
    return [
        Hazard(id=1, image_id=1, object_id="foundation_pit_edge_protection",
               status="confirmed_hazard", hazard_type_id="missing_protection",
               bbox=[0, 278, 999, 999], reasoning_chain=[], visual_evidence="未见连续防护栏杆",
               rule_basis="开挖深度2m及以上未设置防护栏杆", evidence_sufficiency="sufficient",
               confirmed=True),
    ]


def test_report_contains_sections(tmp_path: Path):
    path = build_docx_report(
        session_id=7, hazards=_hazards(), report_dir=str(tmp_path),
        object_name_for=lambda oid: "基坑临边防护",
        hazard_name_for=lambda hid: "防护缺失",
        remediation_for=lambda oid: [{"id": "q1", "condition": "设置连续防护栏杆与挡脚板", "source": "JGJ 80-2016 4.1.2"}],
    )
    assert path.endswith(".docx")
    text = _doc_text(path)
    assert "四口五临边隐患排查报告" in text
    assert "基坑临边防护" in text
    assert "防护缺失" in text
    assert "设置连续防护栏杆与挡脚板" in text
    assert "JGJ 80-2016 4.1.2" in text


def test_empty_report(tmp_path: Path):
    path = build_docx_report(session_id="test_session", hazards=[], report_dir=str(tmp_path),
                             object_name_for=lambda o: o, hazard_name_for=lambda h: h or "",
                             remediation_for=lambda o: [])
    assert "未发现已确认隐患" in _doc_text(path)


def test_report_file_is_valid_docx(tmp_path: Path):
    path = build_docx_report(
        session_id=3, hazards=_hazards(), report_dir=str(tmp_path),
        object_name_for=lambda o: "对象", hazard_name_for=lambda h: "防护缺失",
        remediation_for=lambda o: [],
    )
    doc = Document(path)
    assert len(doc.paragraphs) > 0
