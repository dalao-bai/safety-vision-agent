# backend/tests/test_foe_report_generator.py
import os

from docx import Document

from app.models.foe_schemas import ClauseRef, FoeAnalysis, FoeObject
from app.services import foe_report_generator as foe


class _FakeRetriever:
    def retrieve(self, obj):
        return [ClauseRef(
            standard_code="JGJ 80-2016", clause_id="第4.1.1条",
            official_text="坠落高度基准面2m及以上进行临边作业时，应设置防护栏杆。",
            paraphrase="未设防护", source_raw="x",
        )]


def _analysis():
    return FoeAnalysis(objects=[FoeObject(
        related_object="基坑临边防护", object_bbox=[0, 0, 1, 1],
        status="confirmed_hazard", hazard_type_id="missing_protection",
        hazard_type="防护缺失", visual_evidence="坑边无栏杆",
    )])


def test_generate_report_creates_docx_with_clause(tmp_path):
    path = foe.generate_foe_report_docx(
        [_analysis()], _FakeRetriever(), str(tmp_path), title="测试报告"
    )
    assert os.path.exists(path)
    doc = Document(path)
    full = "\n".join(p.text for p in doc.paragraphs)
    assert "基坑临边防护" in full
    assert "JGJ 80-2016 第4.1.1条" in full
    assert "防护缺失" in full


def test_task_registry_roundtrip():
    tid = foe.create_task("user-1")
    assert foe.get_task(tid)["status"] == "pending"
    foe.set_task_done(tid, "/tmp/r.docx")
    view = foe.get_task(tid)
    assert view["status"] == "done"
    assert view["download_url"] == f"/api/foe/report/download/{tid}"
    assert view["user_id"] == "user-1"


def test_unknown_task_returns_none():
    assert foe.get_task("nope") is None


def test_image_ref_outside_root_not_embedded_and_not_leaked(tmp_path):
    # image_ref 指向 image_root 之外的文件：不嵌入、不泄露路径、不抛异常
    secret = tmp_path / "secret.txt"
    secret.write_text("TOPSECRET")
    analysis = FoeAnalysis(image_ref=str(secret), objects=[FoeObject(
        related_object="基坑临边防护", object_bbox=[0, 0, 1, 1],
        status="safe", visual_evidence="x",
    )])
    out_dir = tmp_path / "out"
    path = foe.generate_foe_report_docx(
        [analysis], _FakeRetriever(), str(out_dir),
        image_root=str(tmp_path / "allowed"),  # secret 不在该目录内
    )
    full = "\n".join(p.text for p in Document(path).paragraphs)
    assert "secret.txt" not in full
    assert str(secret) not in full
    assert "TOPSECRET" not in full


def test_set_task_error_path():
    tid = foe.create_task("u")
    foe.set_task_error(tid, "boom")
    view = foe.get_task(tid)
    assert view["status"] == "error"
    assert view["error"] == "boom"
    assert "download_url" not in view
