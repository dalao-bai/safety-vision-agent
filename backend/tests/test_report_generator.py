"""任务 C 测试：合规报告生成（网络无关）。

用临时 SQLite 造一个用户 + 几条对话/图片/分析结果，调 generate_report_docx，
断言返回路径存在、能被 python-docx 打开、文档里含隐患名称；再覆盖 search_fn
被调用时引用文本进入文档。
"""

from __future__ import annotations

import pytest
from docx import Document

from app.db import repositories as repo
from app.db.sqlite import connect, init_db
from app.services import report_generator as report


@pytest.fixture
def conn(tmp_path):
    connection = connect(str(tmp_path / "test.db"))
    init_db(connection)
    yield connection
    connection.close()


def _seed_user_with_hazards(conn, tmp_path, *, with_image=True):
    """造一个用户 + 一条对话，带一张图片和一条含隐患的分析结果。返回 user_id。"""
    user_id = repo.create_user(conn, "工头老王", "hash")
    cid = repo.create_conversation(conn, user_id=user_id)

    image_id = None
    if with_image:
        # 真造一张可被 python-docx 嵌入的 PNG，确保 add_picture 路径被走到。
        img_path = tmp_path / "hazard.png"
        img_path.write_bytes(_PNG_1PX)
        image_id = repo.add_uploaded_image(
            conn, cid, str(img_path), "hazard.png", "image/png", len(_PNG_1PX)
        )

    repo.save_analysis_result(
        conn,
        cid,
        {
            "summary": "脚手架区域存在多处隐患",
            "hazards": [
                {
                    "name": "未佩戴安全帽",
                    "location": "脚手架二层",
                    "risk_level": "high",
                    "basis": "GB标准要求进入工地必须佩戴安全帽",
                    "remediation": "立即佩戴合格安全帽",
                    "confidence": 0.92,
                }
            ],
            "needs_followup": False,
            "followup_question": None,
        },
        image_id=image_id,
    )
    return user_id


def test_generate_report_creates_openable_docx_with_hazard_name(conn, tmp_path):
    user_id = _seed_user_with_hazards(conn, tmp_path)
    out_dir = tmp_path / "reports"

    path = report.generate_report_docx(
        conn, user_id, "工头老王", None, None, str(out_dir)
    )

    # 路径存在且确实是一份可打开的 docx。
    import os

    assert os.path.exists(path)
    assert path.endswith(".docx")

    document = Document(path)
    full_text = "\n".join(p.text for p in document.paragraphs)
    assert "未佩戴安全帽" in full_text
    # 统计表里的隐患名称（表格文本不在 paragraphs 里，单独收集）。
    table_text = "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    assert "未佩戴安全帽" in table_text


def test_search_fn_reference_text_enters_document(conn, tmp_path):
    user_id = _seed_user_with_hazards(conn, tmp_path, with_image=False)
    out_dir = tmp_path / "reports"

    calls: list[str] = []

    def fake_search(hazard_name: str):
        calls.append(hazard_name)
        return [{"text": "第3.2条：作业人员应正确佩戴安全帽", "source": "安全规范2024"}]

    path = report.generate_report_docx(
        conn, user_id, "工头老王", None, None, str(out_dir), search_fn=fake_search
    )

    # search_fn 确实被调用，且引用原文进入文档。
    assert "未佩戴安全帽" in calls
    document = Document(path)
    full_text = "\n".join(p.text for p in document.paragraphs)
    assert "第3.2条：作业人员应正确佩戴安全帽" in full_text
    assert "安全规范2024" in full_text


def test_empty_range_produces_report_without_hazards(conn, tmp_path):
    user_id = repo.create_user(conn, "空用户", "hash")
    out_dir = tmp_path / "reports"

    path = report.generate_report_docx(
        conn, user_id, "空用户", None, None, str(out_dir)
    )

    document = Document(path)
    full_text = "\n".join(p.text for p in document.paragraphs)
    assert "暂无隐患" in full_text


def test_missing_image_file_is_annotated_not_fatal(conn, tmp_path):
    """图片记录指向不存在的文件时，报告应标注缺失而非抛错。"""
    user_id = repo.create_user(conn, "缺图用户", "hash")
    cid = repo.create_conversation(conn, user_id=user_id)
    image_id = repo.add_uploaded_image(
        conn, cid, str(tmp_path / "gone.png"), "gone.png", "image/png", 10
    )
    repo.save_analysis_result(
        conn,
        cid,
        {
            "summary": "x",
            "hazards": [
                {
                    "name": "临边无防护",
                    "location": "楼层边缘",
                    "risk_level": "critical",
                    "basis": "依据A",
                    "remediation": "加装防护栏",
                    "confidence": 0.8,
                }
            ],
            "needs_followup": False,
            "followup_question": None,
        },
        image_id=image_id,
    )

    path = report.generate_report_docx(
        conn, user_id, "缺图用户", None, None, str(tmp_path / "reports")
    )
    document = Document(path)
    full_text = "\n".join(p.text for p in document.paragraphs)
    assert "临边无防护" in full_text
    assert "图片缺失" in full_text


def test_task_registry_lifecycle_and_ownership():
    task_id = report.create_task("user-1")
    assert report.get_task(task_id)["status"] == "pending"
    assert report.get_task(task_id)["user_id"] == "user-1"

    report.set_task_done(task_id, "/abs/path/report.docx")
    done = report.get_task(task_id)
    assert done["status"] == "done"
    assert done["download_url"] == f"/api/report/download/{task_id}"
    assert done["path"] == "/abs/path/report.docx"

    err_id = report.create_task("user-2")
    report.set_task_error(err_id, "boom")
    err = report.get_task(err_id)
    assert err["status"] == "error"
    assert err["error"] == "boom"

    assert report.get_task("nonexistent") is None


# 最小合法 PNG（1x1 透明像素），供 add_picture 嵌入用。
_PNG_1PX = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00"
    b"\x01\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
)
