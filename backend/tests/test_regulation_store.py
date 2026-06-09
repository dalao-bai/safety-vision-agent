"""RegulationStore 单元测试 (v0.2)。

注入确定性 fake embed_fn（把文本映射成固定维度向量），使用 tmp chroma_dir，
全程不依赖网络。覆盖：chunk_text 切分与 overlap、add+search 召回、delete 后
检索不到、docx 解析。
"""

from __future__ import annotations

import hashlib

import pytest

from app.services.regulation_store import RegulationStore

_DIM = 16


def _fake_embed(texts: list[str]) -> list[list[float]]:
    """确定性嵌入：把文本的关键词哈希成稀疏向量。

    相同/相似词集合的文本距离更近，便于断言召回顺序。每个 token 哈希到一个维度
    并累加权重，向量随后归一化。
    """
    vectors: list[list[float]] = []
    for text in texts:
        vec = [0.0] * _DIM
        for token in text.split():
            h = int(hashlib.md5(token.encode("utf-8")).hexdigest(), 16)
            vec[h % _DIM] += 1.0
        norm = sum(v * v for v in vec) ** 0.5 or 1.0
        vectors.append([v / norm for v in vec])
    return vectors


@pytest.fixture
def store(tmp_path):
    return RegulationStore(chroma_dir=str(tmp_path / "chroma"), embed_fn=_fake_embed)


def test_chunk_text_splits_with_overlap(store):
    text = "a" * 2000
    chunks = store.chunk_text(text, size=800, overlap=100)
    # step = 700 -> 起点 0, 700, 1400 共 3 个 chunk。
    assert len(chunks) == 3
    assert len(chunks[0]) == 800
    # 相邻 chunk 重叠 100 字符：chunk[0] 尾部 == chunk[1] 头部。
    assert chunks[0][-100:] == chunks[1][:100]
    # 末段为剩余部分（2000 - 1400 = 600）。
    assert len(chunks[2]) == 600


def test_chunk_text_empty_returns_empty(store):
    assert store.chunk_text("   ") == []


def test_chunk_text_rejects_bad_overlap(store):
    with pytest.raises(ValueError):
        store.chunk_text("hello", size=100, overlap=100)


def test_add_and_search_recalls_relevant_chunk(store):
    chunks = [
        "scaffold guardrail must be installed at height",
        "fire extinguisher placement requirements indoor",
        "electrical wiring insulation safety inspection",
    ]
    written = store.add_regulation("reg-1", "safety.pdf", chunks)
    assert written == 3

    results = store.search("guardrail height scaffold", top_k=1)
    assert len(results) == 1
    top = results[0]
    assert top["text"] == chunks[0]
    assert top["original_name"] == "safety.pdf"
    assert top["regulation_file_id"] == "reg-1"
    assert isinstance(top["distance"], float)


def test_add_empty_chunks_returns_zero(store):
    assert store.add_regulation("reg-x", "empty.pdf", []) == 0


def test_delete_removes_chunks_from_search(store):
    chunks = ["scaffold guardrail height", "another unrelated topic"]
    store.add_regulation("reg-2", "doc.pdf", chunks)
    assert store.search("scaffold guardrail", top_k=3)

    store.delete_regulation("reg-2")
    assert store.search("scaffold guardrail", top_k=3) == []


def test_delete_only_targets_given_file(store):
    store.add_regulation("reg-a", "a.pdf", ["scaffold guardrail height alpha"])
    store.add_regulation("reg-b", "b.pdf", ["scaffold guardrail height beta"])
    store.delete_regulation("reg-a")

    results = store.search("scaffold guardrail height", top_k=5)
    ids = {r["regulation_file_id"] for r in results}
    assert ids == {"reg-b"}


def test_parse_docx(store, tmp_path):
    from docx import Document

    docx_path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("第一条 施工现场必须佩戴安全帽")
    doc.add_paragraph("")  # 空段落应被忽略
    doc.add_paragraph("第二条 高处作业必须系安全带")
    doc.save(str(docx_path))

    text = store.parse_file(str(docx_path), "docx")
    assert "安全帽" in text
    assert "安全带" in text
    # 空段落被过滤：只剩两行。
    assert len(text.splitlines()) == 2


def test_parse_file_rejects_unsupported_type(store):
    with pytest.raises(ValueError):
        store.parse_file("whatever.txt", "txt")
